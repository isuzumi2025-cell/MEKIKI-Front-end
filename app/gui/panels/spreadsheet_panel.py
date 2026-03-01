import customtkinter as ctk
import tkinter as tk
from tkinter import ttk
from typing import List, Any, Callable, Optional
import concurrent.futures
from PIL import Image, ImageTk
import difflib
import time
from datetime import datetime
from pathlib import Path
from app.gui.managers.thumbnail_manager import ThumbnailManager

LOG_FILE = Path(__file__).parent.parent.parent.parent / "lcs_diagnostic.log"

def log_diagnostic(msg: str):
    """Write a diagnostic message to the panel log file."""
    try:
        timestamp = datetime.now().strftime("%H:%M:%S")
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write(f"[{timestamp}] {msg}\n")
        print(msg)
    except Exception as e:
        print(f"[LOG ERROR] {e}: {msg}")

# #region agent log - Debug instrumentation helper
DEBUG_LOG_PATH = r"c:\Users\raiko\OneDrive\Desktop\26\.cursor\debug.log"
def _debug_log(hypothesis_id: str, location: str, message: str, data: dict = None, run_id: str = "pre-fix"):
    try:
        import json as _json_dbg
        ts = int(time.time() * 1000)
        entry = {
            "hypothesisId": hypothesis_id,
            "location": location,
            "message": message,
            "data": data or {},
            "timestamp": ts,
            "runId": run_id
        }
        with open(DEBUG_LOG_PATH, "a", encoding="utf-8") as f:
            f.write(_json_dbg.dumps(entry, ensure_ascii=False) + "\n")
    except:
        pass
# #endregion

log_diagnostic("=== SpreadsheetPanel MODULE LOADED ===")
log_diagnostic(f"LOG_FILE path: {LOG_FILE}")



class SpreadsheetPanel(ctk.CTkFrame):
    """
    Live Comparison Sheet (Virtual Spreadsheet)
    - Thumbnails below ID (click to jump to Source)
    - Score display with color coding
    - Full text display
    """

    # 螳壽焚
    _simple_widgets_logged = False  # 診断ログ重複抑止

    LCS_FONT_SIZE = 13
    MAX_TEXT_LENGTH = 10000
    PRINT_TARGET_ROWS = 45
    PRINT_TARGET_CHARS = 18000
    THUMB_WIDTH = 60
    THUMB_HEIGHT = 80
    ROW_HEIGHT = 140
    SCORE_WIDTH = 70
    ID_COLUMN_WIDTH = 100  # Web ID/Thumb 縺ｨ PDF ID/Thumb 縺ｮ蟷・ｒ邨ｱ荳
    FULL_RENDER_MAX_ROWS = 180

    def __init__(self, parent, on_row_select: Optional[Callable] = None, **kwargs):
        super().__init__(parent, **kwargs)
        log_diagnostic("[SpreadsheetPanel] __init__ called")

        self.sync_pairs = []
        self._all_sync_pairs = []
        self.web_map = {}
        self.pdf_map = {}
        self.web_image = None
        self.pdf_image = None
        self.selected_indices = None
        self.selected_widget = None
        self.on_row_select = on_row_select
        self._thumbnail_refs = []
        self._export_rows_cache = []

        # Virtual list state
        self._visible_rows = {}  # {index: row_widget}
        self._visible_range = (0, 0)  # (start_index, end_index)
        self._rows_per_page = 20  # Number of visible rows to render
        self._virtual_mode = True  # 軽量仮想描画を既定で安定利用
        self._scroll_update_pending = False
        self._scroll_update_job = None
        self._scroll_refresh_job = None
        self._thumbnail_queue = []
        self._thumbnail_job = None
        self._is_disposed = False
        self._scroll_poll_job = None  # 80ms polling for scroll drag
        self._row_slots = {}  # {index: placeholder_slot_widget}
        self._hydrated_rows = set()  # currently materialized indices
        self._staged_mode = False
        self._staged_buffer_rows = 4
        self._staged_chunk_size = 8
        self._staged_jump_chunk_size = 24
        self._staged_jump_threshold = 40
        self._active_chunk_size = self._staged_chunk_size
        self._thumb_pending_keys = set()
        self._thumb_cache = {}
        self._thumb_cache_order = []
        self._thumb_cache_max = 600
        self._last_scroll_event_ms = 0
        self._last_scroll_yview = 0.0
        # self._executor removed - reverting to safe sync batching

        # scan後固定表示モード（スクロール時の再描画/再計算を停止）
        self._fixed_snapshot_mode = True

        # Re-entry guard and stale thumbnail prevention
        self._render_in_progress = False
        self._pending_visible_range = None
        self._render_epoch = 0

        # 実効行高さ（仮想リスト推定補正用）
        self._row_height_px = self.ROW_HEIGHT + 1

        self._build_ui()
        self.bind("<Destroy>", self._on_destroy, add="+")
        log_diagnostic("[SpreadsheetPanel] UI built successfully")

    def set_on_row_select(self, callback: Callable):
        """Set callback for row selection"""
        self.on_row_select = callback

    def set_images(self, web_image: Image.Image, pdf_image: Image.Image):
        """Set source images for thumbnail generation"""
        self.web_image = web_image
        self.pdf_image = pdf_image

    def _build_ui(self):
        # 1. Toolbar
        toolbar = ctk.CTkFrame(self, height=30, fg_color="#333333")
        toolbar.pack(fill="x", side="top")

        ctk.CTkLabel(toolbar, text="Live Comparison Sheet", font=("Meiryo", 12, "bold")).pack(side="left", padx=10)

        self.stats_label = ctk.CTkLabel(toolbar, text="Web: - | PDF: - | Match: -", font=("Meiryo", 11))
        self.stats_label.pack(side="left", padx=20)

        self.export_btn = ctk.CTkButton(
            toolbar,
            text="Excel Export",
            width=100,
            height=24,
            state="disabled",
            command=self._on_export
        )
        self.export_btn.pack(side="right", padx=5)

        # 2. Table Header
        header_frame = ctk.CTkFrame(self, height=28, fg_color="#2B2B2B")
        header_frame.pack(fill="x", side="top", pady=(1, 0))

        # Header columns - Phase 1.9.7: Actions 繧ｫ繝ｩ繝霑ｽ蜉
        headers = [
            ("Score", 50),
            ("Web ID / Thumb", 100),
            ("Web Text", 0),
            ("", 30),  # Arrow
            ("PDF Text", 0),
            ("PDF ID / Thumb", 100),
            ("Actions", 80),  # 笘・Phase 1.9.7: 譁ｰ隕剰ｿｽ蜉
        ]

        for text, width in headers:
            if width > 0:
                lbl = ctk.CTkLabel(header_frame, text=text, width=width, font=("Meiryo", 9, "bold"))
            else:
                lbl = ctk.CTkLabel(header_frame, text=text, font=("Meiryo", 9, "bold"))
            lbl.pack(side="left", padx=1, pady=4)
            if width == 0:
                lbl.pack_configure(expand=True, fill="x")

        # 3. Scrollable content area
        self.scroll_frame = ctk.CTkScrollableFrame(self, fg_color="#1E1E1E", corner_radius=0)
        self.scroll_frame.pack(fill="both", expand=True)

        # Bind scroll events for virtual list
        self.scroll_frame.bind("<Configure>", self._on_scroll_configure)
        self.scroll_frame._parent_canvas.bind("<MouseWheel>", self._on_scroll_event)
        self.scroll_frame._parent_canvas.bind("<Button-4>", self._on_scroll_event)  # Linux scroll up
        self.scroll_frame._parent_canvas.bind("<Button-5>", self._on_scroll_event)  # Linux scroll down

        self._start_scroll_polling()

    def _start_scroll_polling(self):
        """Start 80ms polling for scroll drag to reliably update visible range."""
        self._stop_scroll_polling()
        if getattr(self, "_virtual_mode", False):
            self._poll_visible_range()

    def _stop_scroll_polling(self):
        """Stop scroll polling (call on destroy)."""
        if getattr(self, "_scroll_poll_job", None):
            try:
                self.after_cancel(self._scroll_poll_job)
            except Exception:
                pass
            self._scroll_poll_job = None

    def _poll_visible_range(self):
        """Poll scroll position every 80ms; update visible range when in virtual mode."""
        if getattr(self, "_fixed_snapshot_mode", False):
            return
        if getattr(self, "_is_disposed", False):
            return
        if not getattr(self, "_virtual_mode", False):
            return  # virtual mode 時のみ継続
        if not self._ensure_ui_ready():
            self._scroll_poll_job = self._safe_after(80, self._poll_visible_range)
            return
        try:
            self._update_visible_range()
        except Exception as e:
            log_diagnostic(f"[ScrollPoll] Error: {e}")
        self._scroll_poll_job = self._safe_after(80, self._poll_visible_range)

    def _normalize_map_key(self, value: Any) -> str:
        """Normalize ID/map keys to avoid whitespace/type mismatches."""
        if value is None:
            return ""
        key = str(value).strip()
        # keep current schema as-is; only normalize obvious formatting noise
        return key

    def _count_matches(self, pairs: List[Any], threshold: float = 0.25) -> int:
        """Single-source match count used by all sheet stats displays."""
        return sum(1 for p in (pairs or []) if float(getattr(p, "similarity", 0.0) or 0.0) >= threshold)

    def _pair_has_display_content(self, pair) -> bool:
        web_id = self._normalize_map_key(getattr(pair, "web_id", "") or "")
        pdf_id = self._normalize_map_key(getattr(pair, "pdf_id", "") or "")
        web_region = self.web_map.get(web_id)
        pdf_region = self.pdf_map.get(pdf_id)
        sim = float(getattr(pair, "similarity", 0.0) or 0.0)
        # 通常ペア: 両側存在 + similarity >= 0.25
        if web_region is not None and pdf_region is not None and sim >= 0.25:
            return True
        # 手動選択ペア: 片側のみでも表示（SEL_ prefix で判定）
        is_manual = (web_id.startswith("SEL") or pdf_id.startswith("SEL"))
        if is_manual and (web_region is not None or pdf_region is not None):
            return True
        return False

    def _ensure_ui_ready(self):
        try:
            if getattr(self, "_is_disposed", False):
                return False
            if not self.winfo_exists():
                return False
            if not hasattr(self, "scroll_frame") or not self.scroll_frame.winfo_exists():
                return False
            if not hasattr(self.scroll_frame, "_parent_canvas") or not self.scroll_frame._parent_canvas.winfo_exists():
                return False
            if not hasattr(self, "stats_label") or not self.stats_label.winfo_exists():
                return False
            return True
        except Exception as e:
            _debug_log(
                "H2",
                "spreadsheet_panel.py:_ensure_ui_ready:error",
                "UI readiness check failed",
                {"error": str(e)}
            )
            return False

    def _safe_after(self, delay_ms: int, callback: Callable):
        """Schedule callback only while this panel is alive."""
        if getattr(self, "_is_disposed", False):
            return None
        try:
            if not self.winfo_exists():
                return None
            return self.after(delay_ms, callback)
        except Exception:
            return None

    def _on_destroy(self, event=None):
        # Ignore child-widget destroy events; handle only this root panel.
        if event is not None and getattr(event, "widget", None) is not self:
            return
        if getattr(self, "_is_disposed", False):
            return
        log_diagnostic(f"[STAGED_INTERRUPT] epoch={self._render_epoch} reason=destroy")
        self._is_disposed = True
        self._cancel_scheduled_jobs()
        self._stop_scroll_polling()

    def update_data(self, sync_pairs: List[Any], web_regions: List[Any], pdf_regions: List[Any],
                    web_image=None, pdf_image=None, web_pages=None, pdf_pages=None):
        """Update data and refresh display
        
        Args:
            web_pages: List of page dicts with 'image' key for page-aware thumbnails
            pdf_pages: List of page dicts with 'image' key for page-aware thumbnails
        """
        if not self._ensure_ui_ready():
            return

        self._all_sync_pairs = list(sync_pairs or [])
        self.sync_pairs = list(self._all_sync_pairs)
        self.web_pages = list(web_pages or [])
        self.pdf_pages = list(pdf_pages or [])

        # 行数に応じて描画モードを切り替え（大量行は段階描画）
        total_rows = len(self.sync_pairs)
        self._virtual_mode = total_rows > self.FULL_RENDER_MAX_ROWS
        self._staged_mode = self._virtual_mode
        self._fixed_snapshot_mode = False
        log_diagnostic(
            f"[_mode] total_rows={total_rows} virtual_mode={self._virtual_mode} staged_mode={self._staged_mode} fixed_snapshot={self._fixed_snapshot_mode}"
        )
        # #region agent log - Hypothesis H1: update_data entry
        _debug_log(
            "H1",
            "spreadsheet_panel.py:update_data:entry",
            "Spreadsheet update_data entry",
            {
                "sync_pairs": len(self.sync_pairs),
                "display_pairs": len(self.sync_pairs),
                "web_regions": len(web_regions),
                "pdf_regions": len(pdf_regions),
                "web_image": list(web_image.size) if web_image else None,
                "pdf_image": list(pdf_image.size) if pdf_image else None,
                "web_pages": len(self.web_pages),
                "pdf_pages": len(self.pdf_pages)
            }
        )
        # #endregion

        # === DIAGNOSTIC LOG START ===
        log_diagnostic("="*60)
        log_diagnostic("[LCS update_data] DIAGNOSTIC")
        log_diagnostic(f"  sync_pairs count: {len(self.sync_pairs)}")
        log_diagnostic(f"  display_pairs count: {len(self.sync_pairs)}")
        log_diagnostic(f"  web_regions count: {len(web_regions)}")
        log_diagnostic(f"  pdf_regions count: {len(pdf_regions)}")
        log_diagnostic(f"  web_image: {web_image.size if web_image else 'None'}")
        log_diagnostic(f"  pdf_image: {pdf_image.size if pdf_image else 'None'}")
        # 笘・PAGE-AWARE險ｺ譁ｭ
        log_diagnostic(f"  web_pages: {len(self.web_pages) if self.web_pages else 'None'}")
        log_diagnostic(f"  pdf_pages: {len(self.pdf_pages) if self.pdf_pages else 'None'}")
        if self.web_pages and len(self.web_pages) > 0:
            first_page = self.web_pages[0]
            log_diagnostic(f"  web_pages[0] image: {first_page.get('image').size if first_page.get('image') else 'None'}")

        if self.sync_pairs:
            sp = self.sync_pairs[0]
            log_diagnostic("  First SyncPair:")
            log_diagnostic(f"    web_id: {sp.web_id}")
            log_diagnostic(f"    pdf_id: {sp.pdf_id}")
            log_diagnostic(f"    web_bbox: {getattr(sp, 'web_bbox', 'N/A')}")
            log_diagnostic(f"    pdf_bbox: {getattr(sp, 'pdf_bbox', 'N/A')}")
            log_diagnostic(f"    web_text: {repr(getattr(sp, 'web_text', 'N/A')[:50]) if getattr(sp, 'web_text', None) else 'None'}")
            log_diagnostic(f"    pdf_text: {repr(getattr(sp, 'pdf_text', 'N/A')[:50]) if getattr(sp, 'pdf_text', None) else 'None'}")

        if web_regions:
            r = web_regions[0]
            log_diagnostic("  First web_region:")
            log_diagnostic(f"    area_code: {getattr(r, 'area_code', 'N/A')}")
            log_diagnostic(f"    id: {getattr(r, 'id', 'N/A')}")
            log_diagnostic(f"    rect: {getattr(r, 'rect', 'N/A')}")
            log_diagnostic(f"    page_id: {getattr(r, 'page_id', 'N/A')}")  # 笘・霑ｽ蜉

        if pdf_regions:
            r = pdf_regions[0]
            log_diagnostic("  First pdf_region:")
            log_diagnostic(f"    area_code: {getattr(r, 'area_code', 'N/A')}")
            log_diagnostic(f"    id: {getattr(r, 'id', 'N/A')}")
            log_diagnostic(f"    rect: {getattr(r, 'rect', 'N/A')}")
        log_diagnostic("="*60)
        # === DIAGNOSTIC LOG END ===

        self.web_map = {}
        self.pdf_map = {}

        for r in web_regions:
            key = getattr(r, 'area_code', None) or getattr(r, 'id', None)
            nkey = self._normalize_map_key(key)
            if nkey:
                self.web_map[nkey] = r

        for r in pdf_regions:
            key = getattr(r, 'area_code', None) or getattr(r, 'id', None)
            nkey = self._normalize_map_key(key)
            if nkey:
                self.pdf_map[nkey] = r

        # Display rows are filtered to avoid huge blank tail (scale mismatch in scrollbar).
        filtered_pairs = [p for p in self._all_sync_pairs if self._pair_has_display_content(p)]
        if filtered_pairs:
            self.sync_pairs = filtered_pairs
        log_diagnostic(f"[DISPLAY_FILTER] raw={len(self._all_sync_pairs)} display={len(self.sync_pairs)}")

        # === KEY VERIFICATION ===
        if self.sync_pairs and self.web_map:
            sp = self.sync_pairs[0]
            found_web = self.web_map.get(self._normalize_map_key(getattr(sp, 'web_id', '')))
            found_pdf = self.pdf_map.get(self._normalize_map_key(getattr(sp, 'pdf_id', '')))
            log_diagnostic(f"[KEY CHECK] web_id={sp.web_id} -> found={found_web is not None}")
            log_diagnostic(f"[KEY CHECK] pdf_id={sp.pdf_id} -> found={found_pdf is not None}")
            if not found_web:
                log_diagnostic(f"[KEY CHECK] web_map keys (first 5): {list(self.web_map.keys())[:5]}")
            if not found_pdf:
                log_diagnostic(f"[KEY CHECK] pdf_map keys (first 5): {list(self.pdf_map.keys())[:5]}")

        if web_image:
            self.web_image = web_image
        if pdf_image:
            self.pdf_image = pdf_image

        # Stats update
        total_web = len(web_regions)
        total_pdf = len(pdf_regions)
        matched = self._count_matches(self._all_sync_pairs, threshold=0.25)
        self.stats_label.configure(text=f"Web: {total_web} | PDF: {total_pdf} | Match: {matched}")

        # Reset scroll to top when applying new data (reduce blank/incomplete display)
        try:
            self.scroll_frame._parent_canvas.yview_moveto(0.0)
        except Exception as e:
            log_diagnostic(f"[update_data] yview reset failed: {e}")

        self._build_export_rows_cache()
        log_diagnostic(f"[fixed_snapshot] rows={len(self._export_rows_cache)} staged={self._staged_mode}")
        self._refresh_rows()

        # #region agent log - Hypothesis H2: update_data exit
        _debug_log(
            "H2",
            "spreadsheet_panel.py:update_data:exit",
            "Spreadsheet update_data exit",
            {
                "visible_rows": len(getattr(self, "_visible_rows", {})),
                "sync_pairs": len(self.sync_pairs)
            }
        )
        # #endregion

        self.export_btn.configure(state="normal" if self.sync_pairs else "disabled")

    def _cancel_scheduled_jobs(self):
        if self._thumbnail_job:
            try:
                self.after_cancel(self._thumbnail_job)
            except Exception:
                pass
            self._thumbnail_job = None

        if self._scroll_update_job:
            try:
                self.after_cancel(self._scroll_update_job)
            except Exception:
                pass
            self._scroll_update_job = None
        self._scroll_refresh_job = None

        self._thumbnail_queue = []
        self._thumb_pending_keys = set()
        self._scroll_update_pending = False

    def _clear_widget_children(self, widget):
        try:
            for child in widget.winfo_children():
                try:
                    if child.winfo_exists():
                        child.destroy()
                except Exception:
                    pass
        except Exception:
            pass

    def _set_slot_placeholder(self, slot, index: int):
        self._clear_widget_children(slot)
        row_bg = "#2D2D2D" if index % 2 == 0 else "#252525"
        slot.configure(fg_color=row_bg)
        ph = tk.Label(slot, text=f"Row {index + 1}", fg="#777777", bg=row_bg, font=("Meiryo", 8))
        ph.pack(anchor="w", padx=8, pady=4)

    def _init_row_slots(self, total_rows: int):
        self._row_slots = {}
        self._hydrated_rows = set()
        for i in range(total_rows):
            row_bg = "#2D2D2D" if i % 2 == 0 else "#252525"
            slot = ctk.CTkFrame(self.scroll_frame, fg_color=row_bg, corner_radius=0, height=self.ROW_HEIGHT)
            slot.pack(fill="x", pady=1)
            slot.pack_propagate(False)
            slot._row_index = i
            self._row_slots[i] = slot
            self._set_slot_placeholder(slot, i)
        self._update_measured_row_height()

    def _update_measured_row_height(self):
        """Measure actual slot height and keep virtual height aligned."""
        try:
            self.scroll_frame.update_idletasks()
            if not self._row_slots:
                return
            first_slot = self._row_slots.get(0)
            if not first_slot or not first_slot.winfo_exists():
                return
            measured = max(1, int(first_slot.winfo_height() or self.ROW_HEIGHT))
            # include inter-row spacing to keep end-of-scroll aligned
            self._row_height_px = measured + 1
        except Exception:
            pass

    def _hydrate_row_slot(self, index: int):
        if index in self._hydrated_rows:
            return
        slot = self._row_slots.get(index)
        if not slot:
            return
        if index >= len(self.sync_pairs):
            return
        self._clear_widget_children(slot)
        pair = self.sync_pairs[index]
        row = self._create_row(index, pair, parent=slot, within_slot=True)
        if row:
            self._visible_rows[index] = row
            self._hydrated_rows.add(index)
        else:
            self._set_slot_placeholder(slot, index)

    def _dehydrate_row_slot(self, index: int):
        if index not in self._hydrated_rows:
            return
        slot = self._row_slots.get(index)
        self._thumbnail_queue = [t for t in self._thumbnail_queue if t.get("row_index") != index]
        self._thumb_pending_keys = {k for k in self._thumb_pending_keys if not str(k).startswith(f"{index}:")}
        if not slot:
            self._hydrated_rows.discard(index)
            self._visible_rows.pop(index, None)
            return
        self._set_slot_placeholder(slot, index)
        self._hydrated_rows.discard(index)
        self._visible_rows.pop(index, None)

    def _make_thumb_key(self, row_index: int, source: str, region) -> str:
        region_id = getattr(region, "area_code", None) or getattr(region, "id", None) or "unknown"
        page_id = getattr(region, "page_id", 0) if region else 0
        rect = tuple(getattr(region, "rect", []) or [])
        return f"{row_index}:{source}:{region_id}:{page_id}:{rect}"

    def _cache_thumb_get(self, key: str):
        thumb = self._thumb_cache.get(key)
        if thumb is None:
            return None
        try:
            self._thumb_cache_order.remove(key)
        except ValueError:
            pass
        self._thumb_cache_order.append(key)
        return thumb

    def _cache_thumb_put(self, key: str, thumb):
        if key in self._thumb_cache:
            self._thumb_cache[key] = thumb
            try:
                self._thumb_cache_order.remove(key)
            except ValueError:
                pass
            self._thumb_cache_order.append(key)
            return
        self._thumb_cache[key] = thumb
        self._thumb_cache_order.append(key)
        while len(self._thumb_cache_order) > self._thumb_cache_max:
            evict_key = self._thumb_cache_order.pop(0)
            self._thumb_cache.pop(evict_key, None)

    def _enqueue_thumbnail_task(self, label, region, source: str, pair, row_index: int):
        if not region:
            return
        key = self._make_thumb_key(row_index, source, region)
        cached = self._cache_thumb_get(key)
        if cached is not None:
            try:
                label.configure(image=cached, text="")
                self._thumbnail_refs.append(cached)
                log_diagnostic(f"[THUMB_CACHE] hit key={key}")
                return
            except Exception:
                pass
        if key in self._thumb_pending_keys:
            log_diagnostic(f"[THUMB_DEDUP] skip key={key}")
            return
        self._thumb_pending_keys.add(key)
        self._thumbnail_queue.append({
            'label': label,
            'region': region,
            'source': source,
            'pair': pair,
            'row_index': row_index,
            'thumb_key': key,
            'epoch': self._render_epoch
        })

    def _render_staged_range(self, start_idx: int, end_idx: int):
        total_rows = len(self.sync_pairs)
        if total_rows == 0:
            return
        keep_start = max(0, start_idx - self._staged_buffer_rows)
        keep_end = min(total_rows, end_idx + self._staged_buffer_rows)
        target_indices = list(range(keep_start, keep_end))

        prev_hydrated = set(self._hydrated_rows)
        for idx in list(prev_hydrated):
            if idx < keep_start or idx >= keep_end:
                self._dehydrate_row_slot(idx)

        hydrate_candidates = [i for i in target_indices if i not in self._hydrated_rows]
        chunk_size = max(1, int(getattr(self, "_active_chunk_size", self._staged_chunk_size)))
        chunk = hydrate_candidates[:chunk_size]
        for idx in chunk:
            self._hydrate_row_slot(idx)

        log_diagnostic(
            f"[STAGED_CHUNK] epoch={self._render_epoch} range=({start_idx},{end_idx}) "
            f"keep=({keep_start},{keep_end}) hydrated_now={len(chunk)} hydrated_total={len(self._hydrated_rows)} "
            f"chunk_size={chunk_size}"
        )
        if len(chunk) < len(hydrate_candidates):
            self._safe_after(1, self._render_visible_rows)

    def _refresh_rows(self):
        """Refresh rows for current render mode (virtual/full)."""
        if self._row_slots or self._hydrated_rows:
            log_diagnostic(f"[STAGED_INTERRUPT] epoch={self._render_epoch} reason=refresh_replace")
        self._cancel_scheduled_jobs()
        self._visible_rows = {}
        self._thumbnail_refs = []
        self._row_slots = {}
        self._hydrated_rows = set()
        self._active_chunk_size = self._staged_chunk_size
        try:
            if not self.scroll_frame.winfo_exists():
                return
            self._clear_widget_children(self.scroll_frame)
        except Exception as e:
            log_diagnostic(f"[_refresh_rows] widget clear failed: {e}")
            return
        total = len(self.sync_pairs)
        if total == 0:
            self._stop_scroll_polling()
            log_diagnostic("[_refresh_rows] rendered 0/0")
            return

        if getattr(self, "_virtual_mode", False):
            self._render_epoch += 1
            try:
                canvas = self.scroll_frame._parent_canvas
                canvas_w = max(1, int(canvas.winfo_width() or 1))
                canvas.configure(scrollregion=(0, 0, canvas_w, max(1, total * max(1, int(self._row_height_px)))))
            except Exception as e:
                log_diagnostic(f"[_refresh_rows] virtual scrollregion init failed: {e}")
            self._init_row_slots(total)
            try:
                self.scroll_frame.update_idletasks()
                canvas = self.scroll_frame._parent_canvas
                bbox = canvas.bbox("all")
                if bbox:
                    canvas.configure(scrollregion=bbox)
                    log_diagnostic(f"[_refresh_rows] staged_scrollregion_bbox={bbox}")
            except Exception as e:
                log_diagnostic(f"[_refresh_rows] staged_scrollregion bbox failed: {e}")
            log_diagnostic(
                f"[STAGED_START] epoch={self._render_epoch} total={total} chunk_size={self._staged_chunk_size}"
            )
            # Force first visible-window render after data replacement.
            self._visible_range = (-1, -1)
            self._start_scroll_polling()
            self._safe_after(0, self._update_visible_range)
            log_diagnostic(f"[_refresh_rows] VIRTUAL INIT total={total}")
            return

        self._stop_scroll_polling()
        self._render_epoch += 1
        for i, pair in enumerate(self.sync_pairs):
            row = self._create_row(i, pair)
            if row:
                self._visible_rows[i] = row
        rendered = len(self._visible_rows)
        self.scroll_frame.update_idletasks()
        try:
            canvas = self.scroll_frame._parent_canvas
            row_h = max(1, int(self._row_height_px))
            total_height = max(row_h, total * row_h)
            canvas_w = max(1, int(canvas.winfo_width() or 1))
            canvas.configure(scrollregion=(0, 0, canvas_w, total_height))
            canvas.yview_moveto(0.0)
            log_diagnostic(f"[_refresh_rows] stable_scrollregion=({canvas_w},{total_height})")
        except Exception as e:
            log_diagnostic(f"[_refresh_rows] scrollregion set failed: {e}")
        self._process_thumbnail_queue()
        log_diagnostic(f"[_refresh_rows] FIXED ALL ROWS: rendered {rendered}/{total}")
        log_diagnostic(f"[STAGED_END] epoch={self._render_epoch} total={total} success={rendered} failed={max(0, total-rendered)}")

    def _render_visible_rows(self):
        """Render only the rows in the visible range"""
        # Re-entry guard: save pending range and defer
        if self._render_in_progress:
            self._pending_visible_range = self._visible_range
            return

        # Guard against TclErrors from already-destroyed widgets.
        try:
            if not self.scroll_frame.winfo_exists():
                log_diagnostic("[_render_visible_rows] scroll_frame does not exist, skipping")
                return
        except Exception as e:
            log_diagnostic(f"[_render_visible_rows] Widget check failed: {e}")
            return

        self._render_in_progress = True

        try:
            start_idx, end_idx = self._visible_range
            if getattr(self, "_staged_mode", False):
                self._render_staged_range(start_idx, end_idx)
                self._process_thumbnail_queue()
                log_diagnostic(
                    f"[STAGED_END] epoch={self._render_epoch} total={len(self.sync_pairs)} "
                    f"success={len(self._hydrated_rows)} failed=0"
                )
                return
            self._render_epoch += 1

            # Clear existing visible rows
            for widget in self.scroll_frame.winfo_children():
                try:
                    if widget.winfo_exists():
                        widget.destroy()
                except Exception:
                    pass  # Widget already destroyed
            self._visible_rows = {}
            self._thumbnail_refs = []
            
            self._cancel_scheduled_jobs()

            # Add top spacer for scrolled-past rows
            if start_idx > 0:
                top_spacer_height = start_idx * max(1, int(self._row_height_px))
                top_spacer = ctk.CTkFrame(self.scroll_frame, height=top_spacer_height, fg_color="#1E1E1E")
                top_spacer.pack(fill="x")
                top_spacer.pack_propagate(False)

            # Render visible rows (耐障害性: 1行失敗で全体停止しない)
            for i in range(start_idx, end_idx):
                if i < len(self.sync_pairs):
                    try:
                        pair = self.sync_pairs[i]
                        row_widget = self._create_row(i, pair)
                        if row_widget:
                            self._visible_rows[i] = row_widget
                    except Exception as e:
                        log_diagnostic(f"[_render_visible_rows] Row {i} failed: {e}")

            rendered_count = len(self._visible_rows)
            if rendered_count == 0 and len(self.sync_pairs) > 0:
                log_diagnostic("[_render_visible_rows] Fallback: no visible rows, rendering safety window")
                fallback_start = max(0, min(start_idx, len(self.sync_pairs) - 1))
                fallback_end = min(len(self.sync_pairs), fallback_start + 12)
                for i in range(fallback_start, fallback_end):
                    try:
                        pair = self.sync_pairs[i]
                        row_widget = self._create_row(i, pair)
                        if row_widget:
                            self._visible_rows[i] = row_widget
                    except Exception as e:
                        log_diagnostic(f"[_render_visible_rows] Fallback row {i} failed: {e}")

            # Add bottom spacer for remaining rows
            remaining_rows = len(self.sync_pairs) - end_idx
            if remaining_rows > 0:
                bottom_spacer_height = remaining_rows * max(1, int(self._row_height_px))
                bottom_spacer = ctk.CTkFrame(self.scroll_frame, height=bottom_spacer_height, fg_color="#1E1E1E")
                bottom_spacer.pack(fill="x")
                bottom_spacer.pack_propagate(False)

            # Ensure geometry is updated before reading yview.
            try:
                self.scroll_frame.update_idletasks()
            except Exception as e:
                log_diagnostic(f"[_render_visible_rows] update_idletasks failed: {e}")

            total_height = 0
            yview_val = (0.0, 0.0)
            try:
                canvas = self.scroll_frame._parent_canvas
                row_h = max(1, int(self._row_height_px))
                total_height = max(row_h, len(self.sync_pairs) * row_h)
                canvas_w = max(1, int(canvas.winfo_width() or 1))
                # Keep scrollbar ratio proportional to total virtual rows.
                canvas.configure(scrollregion=(0, 0, canvas_w, total_height))
                yview_val = canvas.yview()
            except Exception as e:
                log_diagnostic(f"[_render_visible_rows] scrollregion update failed: {e}")

            log_diagnostic(
                f"[_render_visible_rows] yview={yview_val} range=({start_idx},{end_idx}) "
                f"total={len(self.sync_pairs)} visible_widgets={len(self._visible_rows)} total_height={total_height}"
            )

            # Start Async Thumbnail Generation
            self._process_thumbnail_queue()
        finally:
            self._render_in_progress = False
            pending = self._pending_visible_range
            self._pending_visible_range = None
            if pending is not None:
                self._visible_range = pending
                self.after(0, self._render_visible_rows)

    def _process_thumbnail_queue(self):
        """
        Process pending thumbnails in small batches to keep UI responsive.
        Reverted to robust logic: Guarantee processing to verify data integrity.
        """
        self._thumbnail_job = None
        if not self._ensure_ui_ready():
            self._thumbnail_queue = []
            return

        if not self._thumbnail_queue:
            return

        # Process exactly 3 thumbnails per batch (Balanced Speed/Stability)
        # 10ms limit was too aggressive and caused missing thumbnails.
        batch_size = 3
        count = 0
        queue_before = len(self._thumbnail_queue)
        
        while self._thumbnail_queue and count < batch_size:
            task = self._thumbnail_queue.pop(0)
            thumb_key = task.get("thumb_key")
            # Discard stale thumbnails from superseded render
            if task.get('epoch') != self._render_epoch:
                if thumb_key:
                    self._thumb_pending_keys.discard(thumb_key)
                continue
            label = task['label']
            region = task['region']
            source = task['source']
            
            # Skip if label destroyed (scrolled out of view)
            try:
                if not label.winfo_exists():
                    if thumb_key:
                        self._thumb_pending_keys.discard(thumb_key)
                    continue
            except:
                if thumb_key:
                    self._thumb_pending_keys.discard(thumb_key)
                continue

            try:
                # 笘・ByCursor Fix: Page-Aware Thumbnail (page_id繧剃ｽｿ逕ｨ)
                # Generate Thumbnail using page_id for accurate cropping
                pages = getattr(self, 'web_pages', []) if source == "web" else getattr(self, 'pdf_pages', [])
                fallback_image = self.web_image if source == "web" else self.pdf_image
                
                # 蜆ｪ蜈・ Page-Aware (region.page_id縺後≠繧句ｴ蜷・
                if hasattr(region, 'page_id') and region.page_id > 0:
                    thumb = ThumbnailManager.create_page_aware_thumbnail(region, pages, fallback_image)
                else:
                    # 繝輔か繝ｼ繝ｫ繝舌ャ繧ｯ: Global Bbox
                    thumb = self._create_thumbnail(region.rect, source)
                
                if thumb:
                    if thumb_key:
                        self._cache_thumb_put(thumb_key, thumb)
                    self._thumbnail_refs.append(thumb)
                    label.configure(image=thumb, text="") # Remove placeholder text
            except Exception as e:
                print(f"Error generating thumbnail: {e}")
            finally:
                if thumb_key:
                    self._thumb_pending_keys.discard(thumb_key)
            
            count += 1
        
        # Schedule next batch
        log_diagnostic(
            f"[STAGED_THUMB_BATCH] epoch={self._render_epoch} processed={count} "
            f"queue_before={queue_before} queue_remain={len(self._thumbnail_queue)}"
        )
        if self._thumbnail_queue:
            self._thumbnail_job = self._safe_after(20, self._process_thumbnail_queue)

    def _on_scroll_event(self, event):
        """Handle scroll events to trigger virtual list updates"""
        if getattr(self, "_fixed_snapshot_mode", False):
            return
        try:
            self._last_scroll_event_ms = int(time.time() * 1000)
        except Exception:
            self._last_scroll_event_ms = 0
        self._schedule_visible_range_update()
        # Force display refresh after scroll to prevent blank/black on drag (FULL mode)
        if not getattr(self, "_virtual_mode", False):
            if getattr(self, "_scroll_refresh_job", None):
                try: self.after_cancel(self._scroll_refresh_job)
                except Exception: pass
            def _do_refresh():
                self._scroll_refresh_job = None
                if self._ensure_ui_ready():
                    self.scroll_frame.update_idletasks()
            self._scroll_refresh_job = self._safe_after(50, _do_refresh)

    def _on_scroll_configure(self, event):
        """Handle configure events"""
        if getattr(self, "_fixed_snapshot_mode", False):
            return
        self._schedule_visible_range_update()

    def _schedule_visible_range_update(self):
        if getattr(self, "_fixed_snapshot_mode", False):
            return
        if not self._ensure_ui_ready():
            return
        self._scroll_update_pending = True
        if self._scroll_update_job:
            try:
                self.after_cancel(self._scroll_update_job)
            except Exception:
                pass
        now_ms = int(time.time() * 1000)
        elapsed = now_ms - int(getattr(self, "_last_scroll_event_ms", 0) or 0)
        delay_ms = 80
        if 0 <= elapsed < 120:
            delay_ms = 140
        elif 120 <= elapsed < 300:
            delay_ms = 80
        else:
            # after scroll calms down, render quickly
            delay_ms = 20
        self._scroll_update_job = self._safe_after(delay_ms, self._update_visible_range)

    def _update_visible_range(self):
        """Calculate and update visible row range based on scroll position"""
        if getattr(self, "_fixed_snapshot_mode", False):
            return
        self._scroll_update_pending = False
        self._scroll_update_job = None
        self._scroll_refresh_job = None

        if not self.sync_pairs:
            return
        if not getattr(self, "_virtual_mode", False):
            return

        try:
            # Get current scroll position (0.0 to 1.0)
            canvas = self.scroll_frame._parent_canvas
            yview = canvas.yview()
            scroll_top = yview[0]  # Top of visible area (0.0 = top, 1.0 = bottom)
            old_y = float(getattr(self, "_last_scroll_yview", 0.0) or 0.0)
            self._last_scroll_yview = scroll_top

            # Re-sync row height from real slot geometry when available.
            self._update_measured_row_height()

            # Calculate visible row indices
            total_rows = len(self.sync_pairs)
            row_h = max(1, int(self._row_height_px))
            total_height = total_rows * row_h

            # Current scroll position in pixels
            scroll_y_px = max(0.0, float(canvas.canvasy(0)))
            scroll_bottom_px = scroll_y_px + max(1, int(canvas.winfo_height() or 1))

            # Calculate visible row range with fixed buffer and edge snapping
            buffer_rows = int(self._staged_buffer_rows)
            canvas_h = max(1, int(canvas.winfo_height() or 1))
            visible_rows = max(8, int((canvas_h + row_h - 1) / row_h) + 2)
            start_idx = max(0, int(scroll_y_px / row_h) - buffer_rows)
            end_idx = min(total_rows, start_idx + visible_rows + (buffer_rows * 2))

            # Edge snap for precise top/bottom behavior.
            if yview[0] <= 0.005:
                start_idx = 0
                end_idx = min(total_rows, visible_rows + (buffer_rows * 2))

            # Geometry-based bottom detection (more reliable than ratio-only).
            is_bottom_by_geometry = False
            last_slot = self._row_slots.get(total_rows - 1) if total_rows > 0 else None
            if last_slot and last_slot.winfo_exists():
                try:
                    last_bottom = int(last_slot.winfo_y()) + int(last_slot.winfo_height())
                    if scroll_bottom_px >= (last_bottom - max(2, int(row_h * 0.25))):
                        is_bottom_by_geometry = True
                except Exception:
                    pass
            if yview[1] >= 0.995 or is_bottom_by_geometry:
                end_idx = total_rows
                start_idx = max(0, end_idx - (visible_rows + (buffer_rows * 2)))

            # Only update if range changed (即時判定)
            old_start, old_end = self._visible_range
            jump_rows = abs(start_idx - old_start)
            if jump_rows >= self._staged_jump_threshold:
                self._active_chunk_size = self._staged_jump_chunk_size
            else:
                self._active_chunk_size = self._staged_chunk_size
            # Skip tiny jitter updates to reduce queue storms.
            minor_shift = abs(start_idx - old_start) <= 1 and abs(end_idx - old_end) <= 1
            minor_y_delta = abs(scroll_top - old_y) < 0.002
            if (start_idx != old_start or end_idx != old_end) and not (minor_shift and minor_y_delta):
                self._visible_range = (start_idx, end_idx)
                log_diagnostic(
                    f"[_visible_calc] row_h={row_h} canvas_h={canvas_h} start={start_idx} end={end_idx} "
                    f"jump={jump_rows} chunk={self._active_chunk_size} yview={yview} "
                    f"scroll_y_px={int(scroll_y_px)}"
                )
                self._render_visible_rows()

        except Exception as e:
            log_diagnostic(f"[Virtual] Scroll update error: {e}")

    def _create_row(self, index: int, pair, parent=None, within_slot: bool = False):
        """Create a single row with thumbnails below ID"""
        # 診断ログ（一度だけ）
        if not SpreadsheetPanel._simple_widgets_logged:
            log_diagnostic("[_create_row] simple widgets mode active (tk labels/buttons placeholder)")
            SpreadsheetPanel._simple_widgets_logged = True

        # ByCursor Fix: scroll_frame のみ必須、parent_canvas チェックは緩和（黒画面化防止）
        try:
            if not self.scroll_frame.winfo_exists():
                return None
        except Exception:
            return None
            
        row_bg = "#2D2D2D" if index % 2 == 0 else "#252525"

        try:
            row_parent = parent if parent is not None else self.scroll_frame
            row = ctk.CTkFrame(row_parent, fg_color=row_bg, corner_radius=0, height=self.ROW_HEIGHT)
            row._row_index = index
            if within_slot:
                row.pack(fill="both", expand=True)
            else:
                row.pack(fill="x", pady=1)
            row.pack_propagate(False)
        except Exception as e:
            log_diagnostic(f"[Row {index}] Creation failed: {e}")
            return None

        # Get regions from map (string key normalization)
        web_id = self._normalize_map_key(getattr(pair, "web_id", "") or "")
        pdf_id = self._normalize_map_key(getattr(pair, "pdf_id", "") or "")
        web_region = self.web_map.get(web_id)
        pdf_region = self.pdf_map.get(pdf_id)

        # === ROW DIAGNOSTIC (first 3 rows only) ===
        if index < 3:
            log_diagnostic(f"[Row {index}] web_id={web_id}, pdf_id={pdf_id}")
            log_diagnostic(f"  web_region found: {web_region is not None}")
            log_diagnostic(f"  pdf_region found: {pdf_region is not None}")
            log_diagnostic(f"  pair.web_bbox: {getattr(pair, 'web_bbox', None)}")
            log_diagnostic(f"  pair.pdf_bbox: {getattr(pair, 'pdf_bbox', None)}")
            log_diagnostic(f"  self.web_image: {self.web_image.size if self.web_image else 'None'}")
            log_diagnostic(f"  self.pdf_image: {self.pdf_image.size if self.pdf_image else 'None'}")

        # 笘・・笘・SSOT蛻ｷ譁ｰ: region縺九ｉ逶ｴ謗･蜿門ｾ暦ｼ・air邨檎罰縺ｮ髢捺磁蜿ら・繧呈賜髯､・・笘・・笘・        # 繝・く繧ｹ繝・ region縺檎悄螳溘・繧ｽ繝ｼ繧ｹ
        w_txt = web_region.text if web_region and hasattr(web_region, 'text') else ""
        p_txt = pdf_region.text if pdf_region and hasattr(pdf_region, 'text') else ""
        
        web_bbox = web_region.rect if web_region and hasattr(web_region, 'rect') else None
        pdf_bbox = pdf_region.rect if pdf_region and hasattr(pdf_region, 'rect') else None
        
        # === SSOT DIAGNOSTIC ===
        if index < 3:
            log_diagnostic(f"[SSOT] Row {index}: web_region.rect={web_bbox}, text_len={len(w_txt)}")
            log_diagnostic(f"[SSOT] Row {index}: pdf_region.rect={pdf_bbox}, text_len={len(p_txt)}")

        # Score calculation
        sim_percent = int(pair.similarity * 100)
        if sim_percent >= 80:
            score_color, score_bg = "#4CAF50", "#1B3D1B"
        elif sim_percent >= 50:
            score_color, score_bg = "#FFC107", "#3D3D1B"
        elif sim_percent >= 30:
            score_color, score_bg = "#FF9800", "#3D2D1B"
        else:
            score_color, score_bg = "#F44336", "#3D1B1B"

        # === LEGACY PACK ORDER: All LEFT, in sequence ===
        # Score 竊・Web ID 竊・Web Text 竊・Arrow 竊・PDF Text 竊・PDF ID

        # 1. Score (LEFT) - tk.Label で再描画競合回避
        score_frame = ctk.CTkFrame(row, fg_color=score_bg, width=50)
        score_frame.pack(side="left", fill="y", padx=1)
        score_frame.pack_propagate(False)
        score_label = tk.Label(score_frame, text=f"{sim_percent}%", fg=score_color, bg=score_bg,
                              font=("Arial", 10, "bold"))
        score_label.pack(expand=True)

        # 2. Web ID + Thumbnail (LEFT) - tk.Label で再描画競合回避
        web_id_frame = ctk.CTkFrame(row, fg_color=row_bg, width=100)
        web_id_frame.pack(side="left", fill="y", padx=1)
        web_id_frame.pack_propagate(False)

        tk.Label(web_id_frame, text=web_id or "-", fg="#4CAF50", bg=row_bg, font=("Meiryo", 8)).pack(pady=(2, 0))

        # 笘・Async Loading: Placeholder first
        web_thumb_label = tk.Label(web_id_frame, bg=row_bg, cursor="hand2", text="...", fg="#888")
        web_thumb_label.pack(pady=2)
        # Bind click (linkage works even without image)
        web_thumb_label.bind("<Button-1>", lambda e, r=web_region: self._on_thumbnail_click(r, "web", pair))
        
        # Queue for async generation
        if web_region:
            self._enqueue_thumbnail_task(web_thumb_label, web_region, "web", pair, index)

        # 3. Web Text (LEFT, expand)
        web_text_frame = ctk.CTkFrame(row, fg_color=row_bg)
        web_text_frame.pack(side="left", fill="both", expand=True, padx=2)

        web_text_widget = tk.Text(web_text_frame, bg=row_bg, fg="#E0E0E0", relief="flat",
                          font=("Meiryo", 9), wrap="word", height=5, width=25)
        web_text_widget.pack(fill="both", expand=True, padx=2, pady=2)

        # 4. Arrow (LEFT) - tk.Label で再描画競合回避
        tk.Label(row, text="<>", width=4, fg="#666666", bg=row_bg).pack(side="left", padx=1)

        # 5. PDF Text (LEFT, expand)
        pdf_text_frame = ctk.CTkFrame(row, fg_color=row_bg)
        pdf_text_frame.pack(side="left", fill="both", expand=True, padx=2)

        pdf_text_widget = tk.Text(pdf_text_frame, bg=row_bg, fg="#E0E0E0", relief="flat",
                          font=("Meiryo", 9), wrap="word", height=5, width=25)
        pdf_text_widget.pack(fill="both", expand=True, padx=2, pady=2)
        
        # 笘・Diff Highlight 驕ｩ逕ｨ (荳閾ｴ=繧ｰ繝ｬ繝ｼ縲∝ｷｮ蛻・襍､/髱・
        try:
            self._apply_diff_highlight(web_text_widget, pdf_text_widget, w_txt, p_txt)
        except Exception as e:
            log_diagnostic(f"[Row {index}] _apply_diff_highlight failed: {e}")
            try:
                web_text_widget.insert("1.0", (w_txt or "")[:self.MAX_TEXT_LENGTH])
                pdf_text_widget.insert("1.0", (p_txt or "")[:self.MAX_TEXT_LENGTH])
                web_text_widget.configure(state="disabled")
                pdf_text_widget.configure(state="disabled")
            except Exception as e2:
                log_diagnostic(f"[Row {index}] Fallback insert failed: {e2}")

        # 6. PDF ID + Thumbnail (LEFT - last)
        pdf_id_frame = ctk.CTkFrame(row, fg_color=row_bg, width=100)
        pdf_id_frame.pack(side="left", fill="y", padx=1)
        pdf_id_frame.pack_propagate(False)

        tk.Label(pdf_id_frame, text=pdf_id or "-", fg="#2196F3", bg=row_bg, font=("Meiryo", 8)).pack(pady=(2, 0))

        # 笘・Async Loading: Placeholder first
        pdf_thumb_label = tk.Label(pdf_id_frame, bg=row_bg, cursor="hand2", text="...", fg="#888")
        pdf_thumb_label.pack(pady=2)
        # Bind click
        pdf_thumb_label.bind("<Button-1>", lambda e, r=pdf_region: self._on_thumbnail_click(r, "pdf", pair))

        # Queue for async generation
        if pdf_region:
             self._enqueue_thumbnail_task(pdf_thumb_label, pdf_region, "pdf", pair, index)

        # Actions列 - CTkButton 再描画競合回避のためプレースホルダに置換
        action_frame = ctk.CTkFrame(row, fg_color=row_bg, width=80)
        action_frame.pack(side="left", fill="y", padx=2)
        action_frame.pack_propagate(False)
        tk.Label(action_frame, text="-", fg="#888", bg=row_bg).pack(pady=(8, 0))

        # Row click binding (text frame も含める)
        row.bind("<Button-1>", lambda e, p=pair, w=row: self._on_row_click(w, p))
        for widget in [score_frame, web_id_frame, pdf_id_frame, web_text_frame, pdf_text_frame]:
            widget.bind("<Button-1>", lambda e, p=pair, w=row: self._on_row_click(w, p))

        return row

    def _create_page_aware_thumbnail(self, region, pages_list, fallback_image, source: str):
        """Delegate to ThumbnailManager"""
        return ThumbnailManager.create_page_aware_thumbnail(region, pages_list, fallback_image)
    
    def _create_thumbnail(self, bbox, source="web"):
        """Delegate to ThumbnailManager (Virtual Cropping from Global BBox)"""
        pages = getattr(self, 'web_pages', []) if source == "web" else getattr(self, 'pdf_pages', [])
        image = self.web_image if source == "web" else self.pdf_image
        
        return ThumbnailManager.create_thumbnail_from_global_bbox(bbox, pages, image)

    def _on_thumbnail_click(self, region, source: str, pair):
        """Handle thumbnail click - notify parent to highlight region (Canonical pattern)"""
        if self.on_row_select and region:
            # 笘・Use region.area_code for proper Source panel linkage
            if source == "web":
                self.on_row_select(region.area_code, pair.pdf_id, pair)
            else:
                self.on_row_select(pair.web_id, region.area_code, pair)
            log_diagnostic(f"[Thumb] Clicked: {source} - {getattr(region, 'area_code', 'N/A')}")

    def _on_row_click(self, row_widget, pair):
        """Handle row click - highlight and notify parent"""
        web_id = self._normalize_map_key(getattr(pair, "web_id", "") or "")
        pdf_id = self._normalize_map_key(getattr(pair, "pdf_id", "") or "")

        # Reset previous selection
        if self.selected_widget:
            try:
                idx = getattr(self.selected_widget, "_row_index", None)
                if idx is None:
                    idx = list(self.scroll_frame.winfo_children()).index(self.selected_widget)
                color = "#2D2D2D" if idx % 2 == 0 else "#252525"
                self.selected_widget.configure(fg_color=color)
            except:
                pass

        # Highlight new selection
        self.selected_widget = row_widget
        self.selected_indices = (web_id, pdf_id)
        row_widget.configure(fg_color="#444466")

        # Notify parent
        if self.on_row_select:
            self.on_row_select(web_id, pdf_id, pair)

    def get_selected_ids(self):
        """Return (web_id, pdf_id) or None"""
        return self.selected_indices
    
    def set_on_similar_search(self, callback: Callable):
        """Set callback for Similar Search button"""
        self.on_similar_search = callback
    
    def set_on_match_search(self, callback: Callable):
        """Set callback for Match Search button"""
        self.on_match_search = callback
    
    def set_on_sync_recalculate(self, callback: Callable):
        """Set callback for Individual Sync Recalculation button"""
        self.on_sync_recalculate = callback
    
    def _on_similar_search(self, pair):
        """Handle Similar Search button click"""
        log_diagnostic(f"[SimilarSearch] Triggered for pair: web={pair.web_id}, pdf={pair.pdf_id}")
        if hasattr(self, 'on_similar_search') and self.on_similar_search:
            self.on_similar_search(pair)
        else:
            log_diagnostic("[SimilarSearch] No callback set")
            print("[SimilarSearch] callback is not set")
    
    def _on_match_search(self, pair):
        """Handle Match Search button click"""
        log_diagnostic(f"[MatchSearch] Triggered for pair: web={pair.web_id}, pdf={pair.pdf_id}")
        if hasattr(self, 'on_match_search') and self.on_match_search:
            self.on_match_search(pair)
        else:
            log_diagnostic("[MatchSearch] No callback set")
            print("[MatchSearch] callback is not set")
    
    def _on_match_with_sync(self, pair):
        """Handle Match Search button click with individual Sync recalculation"""
        log_diagnostic(f"[MatchWithSync] Triggered for pair: web={pair.web_id}, pdf={pair.pdf_id}")
        
        if hasattr(self, 'on_match_search') and self.on_match_search:
            self.on_match_search(pair)
        
        if hasattr(self, 'on_sync_recalculate') and self.on_sync_recalculate:
            self.on_sync_recalculate(pair)
    
    def _on_individual_sync(self, pair):
        """Handle Individual Sync Recalculation button click"""
        log_diagnostic(f"[IndividualSync] Triggered for pair: web={pair.web_id}, pdf={pair.pdf_id}")
        if hasattr(self, 'on_sync_recalculate') and self.on_sync_recalculate:
            self.on_sync_recalculate(pair)
        else:
            # 繝・ヵ繧ｩ繝ｫ繝亥虚菴・ difflib縺ｧ蜀崎ｨ育ｮ励＠縺ｦUI縺ｫ蜿肴丐
            self._default_sync_recalc(pair)
    
    def _default_sync_recalc(self, pair):
        """Default individual sync recalculation using difflib."""
        try:
            from difflib import SequenceMatcher
            
            web_text = getattr(pair, 'web_text', '')
            pdf_text = getattr(pair, 'pdf_text', '')
            
            if web_text and pdf_text:
                new_score = SequenceMatcher(None, web_text, pdf_text).ratio()
                pair.similarity = new_score
                log_diagnostic(f"[IndividualSync] Recalculated: {new_score:.0%}")
                print(f"筺ｳ 蜀崎ｨ育ｮ怜ｮ御ｺ・ {pair.web_id} 竊・{pair.pdf_id} = {new_score:.0%}")
                
                # UI繧偵Μ繝輔Ξ繝・す繝･
                self._refresh_rows()
            else:
                print("筺ｳ 蜀崎ｨ育ｮ・ 繝・く繧ｹ繝医′縺ゅｊ縺ｾ縺帙ｓ")
        except Exception as e:
            log_diagnostic(f"[IndividualSync] Error: {e}")
            print(f"筺ｳ 蜀崎ｨ育ｮ励お繝ｩ繝ｼ: {e}")
    
    def _apply_diff_highlight(self, w_widget, p_widget, t1: str, t2: str):
        """繝・く繧ｹ繝亥ｷｮ蛻・ｒ濶ｲ蛻・￠縺励※陦ｨ遉ｺ (謾ｹ濶ｯ迚・ 繧ｷ繝ｳ繝励Ν縺ｧ豁｣遒ｺ)
        
        繝ｫ繝ｼ繝ｫ:
        - 荳｡譁ｹ縺ｫ蟄伜惠縺吶ｋ繝・く繧ｹ繝・竊・逋ｽ (normal)
        - 迚・婿縺ｮ縺ｿ縺ｮ繝・く繧ｹ繝・竊・邱・(diff)
        """
        import re
        
        w_widget.tag_config("normal", foreground="#FFFFFF")
        w_widget.tag_config("diff", foreground="#4CAF50", background="#1A3D1A")  # 蟾ｮ蛻・邱題レ譎ｯ
        p_widget.tag_config("normal", foreground="#FFFFFF")
        p_widget.tag_config("diff", foreground="#4CAF50", background="#1A3D1A")
        
        if not t1 and not t2:
            w_widget.configure(state="disabled")
            p_widget.configure(state="disabled")
            return
        
        def clean(text):
            if not text:
                return ""
            return re.sub(r'\s+', '', text)[:self.MAX_TEXT_LENGTH]
        
        clean_t1 = clean(t1)
        clean_t2 = clean(t2)
        
        # SequenceMatcher縺ｧ豁｣遒ｺ縺ｪ繝槭ャ繝√Φ繧ｰ
        matcher = difflib.SequenceMatcher(None, clean_t1, clean_t2)
        
        t1_match_ranges = []
        t2_match_ranges = []
        
        for block in matcher.get_matching_blocks():
            if block.size >= 2:  # 2譁・ｭ嶺ｻ･荳翫・荳閾ｴ繧呈治逕ｨ
                t1_match_ranges.append((block.a, block.a + block.size))
                t2_match_ranges.append((block.b, block.b + block.size))
        
        # t1繧定牡蛻・￠縺励※陦ｨ遉ｺ
        pos = 0
        for start, end in t1_match_ranges:
            if pos < start:
                w_widget.insert("end", clean_t1[pos:start], "diff")
            # 繝槭ャ繝・Κ蛻・竊・逋ｽ
            w_widget.insert("end", clean_t1[start:end], "normal")
            pos = end
        # 谿九ｊ
        if pos < len(clean_t1):
            w_widget.insert("end", clean_t1[pos:], "diff")
        
        # t2繧定牡蛻・￠縺励※陦ｨ遉ｺ
        pos = 0
        for start, end in t2_match_ranges:
            if pos < start:
                p_widget.insert("end", clean_t2[pos:start], "diff")
            p_widget.insert("end", clean_t2[start:end], "normal")
            pos = end
        if pos < len(clean_t2):
            p_widget.insert("end", clean_t2[pos:], "diff")
        
        w_widget.configure(state="disabled")
        p_widget.configure(state="disabled")

    def _build_export_rows_cache(self):
        """
        sync_pairs 全件を走査し、export用の行dictリストを構築。
        update_data完了時点で呼び、スクロール描画に依存しないexportスナップショットを確定する。
        """
        rows = []
        source_pairs = getattr(self, "_all_sync_pairs", None) or self.sync_pairs or []
        for i, pair in enumerate(source_pairs):
            try:
                web_id = self._normalize_map_key(getattr(pair, "web_id", "") or "")
                pdf_id = self._normalize_map_key(getattr(pair, "pdf_id", "") or "")
                web_region = self.web_map.get(web_id)
                pdf_region = self.pdf_map.get(pdf_id)
                w_text = getattr(pair, "web_text", "") or (web_region.text if web_region and hasattr(web_region, "text") else "")
                p_text = getattr(pair, "pdf_text", "") or (pdf_region.text if pdf_region and hasattr(pdf_region, "text") else "")
                w_text = (w_text or "")[:500]
                p_text = (p_text or "")[:500]
                sim = float(getattr(pair, "similarity", 0.0) or 0.0)
                sync_percent = int(sim * 100)
                rows.append({
                    "no": i + 1,
                    "web_id": getattr(pair, "web_id", "") or "",
                    "pdf_id": getattr(pair, "pdf_id", "") or "",
                    "web_text": w_text,
                    "pdf_text": p_text,
                    "sync_percent": sync_percent,
                })
            except Exception as e:
                log_diagnostic(f"[_build_export_rows_cache] row {i} failed: {e}")
        self._export_rows_cache = rows
        log_diagnostic(f"[_build_export_rows_cache] cached {len(rows)} rows")

    def _chunk_pairs_for_print(self, pairs, max_rows=45, max_chars=18000):
        """
        行数 + text量(web_text + pdf_text + 少量バッファ)でチャンク分割する。
        印刷用に最適なチャンクを返す。
        """
        buffer_per_row = 50  # 区切り・ヘッダー等のバッファ
        chunks = []
        current_chunk = []
        current_chars = 0

        for pair in (pairs or []):
            web_region = self.web_map.get(self._normalize_map_key(getattr(pair, "web_id", "") or ""))
            pdf_region = self.pdf_map.get(self._normalize_map_key(getattr(pair, "pdf_id", "") or ""))
            w_text = getattr(pair, "web_text", "") or (web_region.text[:500] if web_region else "")
            p_text = getattr(pair, "pdf_text", "") or (pdf_region.text[:500] if pdf_region else "")
            row_chars = len(str(w_text)) + len(str(p_text)) + buffer_per_row

            if current_chunk and (
                len(current_chunk) >= max_rows or (current_chars + row_chars) >= max_chars
            ):
                chunks.append((list(current_chunk), current_chars))
                current_chunk = []
                current_chars = 0

            current_chunk.append(pair)
            current_chars += row_chars

        if current_chunk:
            chunks.append((current_chunk, current_chars))
        return chunks

    def _chunk_export_rows_for_print(self, rows, max_rows=45, max_chars=18000):
        """
        export用キャッシュ行（dict）をチャンク分割する。
        row_chars = len(web_text)+len(pdf_text)+buffer で計算。
        """
        buffer_per_row = 50
        chunks = []
        current_chunk = []
        current_chars = 0
        for row in (rows or []):
            w_text = row.get("web_text", "") or ""
            p_text = row.get("pdf_text", "") or ""
            row_chars = len(str(w_text)) + len(str(p_text)) + buffer_per_row
            if current_chunk and (
                len(current_chunk) >= max_rows or (current_chars + row_chars) >= max_chars
            ):
                chunks.append((list(current_chunk), current_chars))
                current_chunk = []
                current_chars = 0
            current_chunk.append(row)
            current_chars += row_chars
        if current_chunk:
            chunks.append((current_chunk, current_chars))
        return chunks

    def _on_export(self):
        """Export to Excel (印刷用チャンク分割: Print_01, Print_02, ... + Print_Index)
        キャッシュ(_export_rows_cache)を使用。スクロール描画に依存しない。
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
            from pathlib import Path
            from datetime import datetime

            rows = getattr(self, "_export_rows_cache", None) or []
            wb = openpyxl.Workbook()
            chunks = self._chunk_export_rows_for_print(
                rows,
                max_rows=self.PRINT_TARGET_ROWS,
                max_chars=self.PRINT_TARGET_CHARS,
            )

            # Print_Index を先頭に作成
            idx_ws = wb.active
            idx_ws.title = "Print_Index"
            idx_headers = ["Sheet", "Rows", "EstimatedChars", "ExportedAt"]
            for col, h in enumerate(idx_headers, 1):
                c = idx_ws.cell(row=1, column=col, value=h)
                c.font = Font(bold=True)
                c.fill = PatternFill(start_color="4CAF50", fill_type="solid")
            now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            for row_idx, (chunk_rows, est_chars) in enumerate(chunks, 2):
                sheet_name = f"Print_{row_idx - 1:02d}"
                idx_ws.cell(row=row_idx, column=1, value=sheet_name)
                idx_ws.cell(row=row_idx, column=2, value=len(chunk_rows))
                idx_ws.cell(row=row_idx, column=3, value=est_chars)
                idx_ws.cell(row=row_idx, column=4, value=now_str)
            idx_ws.column_dimensions["A"].width = 12
            idx_ws.column_dimensions["B"].width = 8
            idx_ws.column_dimensions["C"].width = 16
            idx_ws.column_dimensions["D"].width = 22

            headers = ["No", "Web ID", "Web Text", "PDF ID", "PDF Text", "Sync %"]
            for chunk_idx, (chunk_rows, _) in enumerate(chunks, 1):
                sheet_name = f"Print_{chunk_idx:02d}"
                ws = wb.create_sheet(title=sheet_name)
                for col, h in enumerate(headers, 1):
                    c = ws.cell(row=1, column=col, value=h)
                    c.font = Font(bold=True)
                    c.fill = PatternFill(start_color="4CAF50", fill_type="solid")

                for i, r in enumerate(chunk_rows, 2):
                    ws.cell(row=i, column=1, value=r.get("no", i - 1))
                    ws.cell(row=i, column=2, value=r.get("web_id", ""))
                    ws.cell(row=i, column=3, value=r.get("web_text", ""))
                    ws.cell(row=i, column=4, value=r.get("pdf_id", ""))
                    ws.cell(row=i, column=5, value=r.get("pdf_text", ""))
                    ws.cell(row=i, column=6, value=f"{r.get('sync_percent', 0)}%")

                ws.freeze_panes = "A2"
                ws.print_title_rows = "1:1"
                ws.page_setup.orientation = "landscape"
                ws.page_setup.fitToWidth = 1
                ws.page_setup.fitToHeight = 0
                ws.column_dimensions["A"].width = 6
                ws.column_dimensions["B"].width = 12
                ws.column_dimensions["C"].width = 60
                ws.column_dimensions["D"].width = 12
                ws.column_dimensions["E"].width = 60
                ws.column_dimensions["F"].width = 10

            Path("./exports").mkdir(exist_ok=True)
            filename = f"comparison_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            output_path = f"./exports/{filename}"
            wb.save(output_path)

            print(f"Excel exported: {output_path}")
            import os
            os.startfile(output_path)

        except Exception as e:
            print(f"Export error: {e}")
            import tkinter.messagebox as mb
            mb.showerror("Export Error", str(e))


    def destroy(self):
        """Ensure pending jobs are stopped before widget teardown."""
        self._on_destroy()
        try:
            super().destroy()
        except Exception:
            pass
