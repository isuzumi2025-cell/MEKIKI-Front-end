"""
Layout Exporter - 多形式エクスポート
Phase 3.0 SDK

Supports:
- PowerPoint (.pptx) - 座標配置再現
- Excel (.xlsx) - レイヤークラスター別シート
- CSV - メタデータ出力
- PNG - 個別画像出力
"""
import os
import csv
from pathlib import Path
from typing import Dict, List, Optional, Any
from PIL import Image


class LayoutExporter:
    """
    多形式エクスポートクラス
    
    レイヤー結果を各種形式にエクスポート。
    """
    
    def __init__(self, layer_result: Any = None, page_images: List[Image.Image] = None):
        """
        Args:
            layer_result: LayerSeparator.extract_layers()の結果
            page_images: PDFページ画像リスト
        """
        self.layer_result = layer_result
        self.page_images = page_images or []
        self._paragraphs: List[Any] = []
    
    def set_data(self, layer_result: Any, page_images: List[Image.Image] = None, paragraphs: List[Any] = None):
        """データを設定"""
        self.layer_result = layer_result
        if page_images:
            self.page_images = page_images
        if paragraphs:
            self._paragraphs = paragraphs
    
    def to_pptx(self, output_path: str, dpi: float = 96.0) -> bool:
        """
        PowerPointにエクスポート（座標配置再現）
        
        Args:
            output_path: 出力パス
            dpi: 座標変換用DPI
        
        Returns:
            成功時True
        """
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RgbColor
            
            prs = Presentation()
            
            # スライドサイズを16:9に設定
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
            
            # 空白レイアウト
            blank_layout = prs.slide_layouts[6]
            
            # ページごとにスライド作成
            if self.page_images:
                for page_num, page_img in enumerate(self.page_images):
                    slide = prs.slides.add_slide(blank_layout)
                    
                    # 背景画像を追加
                    img_path = f"_temp_page_{page_num}.png"
                    page_img.save(img_path)
                    slide.shapes.add_picture(img_path, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
                    os.remove(img_path)
                    
                    # テキストブロックを配置
                    self._add_text_blocks_to_slide(slide, page_num, dpi)
            
            prs.save(output_path)
            print(f"✅ PowerPoint exported: {output_path}")
            return True
            
        except ImportError:
            print("⚠️ python-pptx not installed. Run: pip install python-pptx")
            return False
        except Exception as e:
            print(f"❌ PowerPoint export error: {e}")
            return False
    
    def _add_text_blocks_to_slide(self, slide, page_num: int, dpi: float):
        """スライドにテキストブロックを追加"""
        from pptx.util import Inches, Pt
        
        if not self.layer_result:
            return
        
        text_layer = getattr(self.layer_result, 'text_layer', None)
        if not text_layer or not hasattr(text_layer, 'blocks'):
            return
        
        for block in text_layer.blocks:
            if getattr(block, 'page_num', 0) != page_num:
                continue
            
            bbox = getattr(block, 'bbox', None)
            if not bbox or len(bbox) < 4:
                continue
            
            text = getattr(block, 'text', '')
            if not text:
                continue
            
            # 座標変換 (pixel to inches)
            x0, y0, x1, y1 = bbox[:4]
            left = Inches(x0 / dpi)
            top = Inches(y0 / dpi)
            width = Inches((x1 - x0) / dpi)
            height = Inches((y1 - y0) / dpi)
            
            # テキストボックス追加
            textbox = slide.shapes.add_textbox(left, top, width, height)
            tf = textbox.text_frame
            tf.word_wrap = True
            
            p = tf.paragraphs[0]
            p.text = text
            p.font.size = Pt(10)
    
    def to_excel(self, output_path: str) -> bool:
        """
        Excelにエクスポート（レイヤークラスター別シート）
        
        Args:
            output_path: 出力パス
        
        Returns:
            成功時True
        """
        try:
            from openpyxl import Workbook
            from openpyxl.drawing.image import Image as XLImage
            import io
            
            wb = Workbook()
            
            # デフォルトシートを削除
            if "Sheet" in wb.sheetnames:
                del wb["Sheet"]
            
            # レイヤー別シート作成
            self._create_text_sheets(wb)
            self._create_image_sheet(wb)
            
            wb.save(output_path)
            print(f"✅ Excel exported: {output_path}")
            return True
            
        except ImportError:
            print("⚠️ openpyxl not installed. Run: pip install openpyxl")
            return False
        except Exception as e:
            print(f"❌ Excel export error: {e}")
            return False
    
    def _create_text_sheets(self, wb):
        """テキストレイヤーシートを作成"""
        from openpyxl.utils import get_column_letter
        
        if not self.layer_result:
            return
        
        text_layer = getattr(self.layer_result, 'text_layer', None)
        if not text_layer or not hasattr(text_layer, 'blocks'):
            return
        
        # レイヤー別に分類
        layers = {"Bold": [], "Regular": [], "Light": []}
        
        for block in text_layer.blocks:
            font_size = getattr(block, 'font_size', 12.0)
            if font_size >= 18:
                layers["Bold"].append(block)
            elif font_size >= 12:
                layers["Regular"].append(block)
            else:
                layers["Light"].append(block)
        
        # 各レイヤーのシート作成
        for layer_name, blocks in layers.items():
            if not blocks:
                continue
            
            ws = wb.create_sheet(title=f"Text_{layer_name}")
            
            # ヘッダー
            headers = ["#", "Page", "Text", "X", "Y", "Width", "Height"]
            for col, header in enumerate(headers, 1):
                ws.cell(row=1, column=col, value=header)
            
            # データ
            for i, block in enumerate(blocks, 1):
                bbox = getattr(block, 'bbox', (0, 0, 0, 0))
                ws.cell(row=i+1, column=1, value=i)
                ws.cell(row=i+1, column=2, value=getattr(block, 'page_num', 0) + 1)
                ws.cell(row=i+1, column=3, value=getattr(block, 'text', ''))
                ws.cell(row=i+1, column=4, value=bbox[0] if len(bbox) > 0 else 0)
                ws.cell(row=i+1, column=5, value=bbox[1] if len(bbox) > 1 else 0)
                ws.cell(row=i+1, column=6, value=bbox[2] - bbox[0] if len(bbox) > 2 else 0)
                ws.cell(row=i+1, column=7, value=bbox[3] - bbox[1] if len(bbox) > 3 else 0)
            
            # 列幅調整
            ws.column_dimensions['C'].width = 50
    
    def _create_image_sheet(self, wb):
        """画像レイヤーシートを作成"""
        if not self.layer_result:
            return
        
        image_layer = getattr(self.layer_result, 'image_layer', None)
        if not image_layer or not hasattr(image_layer, 'images'):
            return
        
        ws = wb.create_sheet(title="Images")
        
        # ヘッダー
        headers = ["#", "Page", "X", "Y", "Width", "Height"]
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)
        
        # データ
        for i, img_block in enumerate(image_layer.images, 1):
            bbox = getattr(img_block, 'bbox', (0, 0, 0, 0))
            ws.cell(row=i+1, column=1, value=i)
            ws.cell(row=i+1, column=2, value=getattr(img_block, 'page_num', 0) + 1)
            ws.cell(row=i+1, column=3, value=bbox[0] if len(bbox) > 0 else 0)
            ws.cell(row=i+1, column=4, value=bbox[1] if len(bbox) > 1 else 0)
            ws.cell(row=i+1, column=5, value=bbox[2] - bbox[0] if len(bbox) > 2 else 0)
            ws.cell(row=i+1, column=6, value=bbox[3] - bbox[1] if len(bbox) > 3 else 0)
    
    def to_csv(self, output_path: str) -> bool:
        """
        CSVにエクスポート
        
        Args:
            output_path: 出力パス
        
        Returns:
            成功時True
        """
        try:
            rows = []
            
            # テキストブロック
            if self.layer_result:
                text_layer = getattr(self.layer_result, 'text_layer', None)
                if text_layer and hasattr(text_layer, 'blocks'):
                    for i, block in enumerate(text_layer.blocks, 1):
                        bbox = getattr(block, 'bbox', (0, 0, 0, 0))
                        font_size = getattr(block, 'font_size', 12.0)
                        layer = "Bold" if font_size >= 18 else ("Regular" if font_size >= 12 else "Light")
                        
                        rows.append({
                            "id": i,
                            "type": "text",
                            "layer": layer,
                            "text": getattr(block, 'text', ''),
                            "x": bbox[0] if len(bbox) > 0 else 0,
                            "y": bbox[1] if len(bbox) > 1 else 0,
                            "width": bbox[2] - bbox[0] if len(bbox) > 2 else 0,
                            "height": bbox[3] - bbox[1] if len(bbox) > 3 else 0,
                            "page": getattr(block, 'page_num', 0) + 1
                        })
                
                # 画像
                image_layer = getattr(self.layer_result, 'image_layer', None)
                if image_layer and hasattr(image_layer, 'images'):
                    for i, img_block in enumerate(image_layer.images):
                        bbox = getattr(img_block, 'bbox', (0, 0, 0, 0))
                        rows.append({
                            "id": len(rows) + 1,
                            "type": "image",
                            "layer": f"Image{i+1}",
                            "text": "",
                            "x": bbox[0] if len(bbox) > 0 else 0,
                            "y": bbox[1] if len(bbox) > 1 else 0,
                            "width": bbox[2] - bbox[0] if len(bbox) > 2 else 0,
                            "height": bbox[3] - bbox[1] if len(bbox) > 3 else 0,
                            "page": getattr(img_block, 'page_num', 0) + 1
                        })
            
            # CSV書き出し
            if rows:
                with open(output_path, 'w', newline='', encoding='utf-8-sig') as f:
                    writer = csv.DictWriter(f, fieldnames=rows[0].keys())
                    writer.writeheader()
                    writer.writerows(rows)
            
            print(f"✅ CSV exported: {output_path}")
            return True
            
        except Exception as e:
            print(f"❌ CSV export error: {e}")
            return False
    
    def to_png(self, output_dir: str) -> bool:
        """
        個別PNG画像として出力
        
        Args:
            output_dir: 出力ディレクトリ
        
        Returns:
            成功時True
        """
        try:
            Path(output_dir).mkdir(parents=True, exist_ok=True)
            exported = 0
            
            # テキストブロックのクロップ画像
            if self.layer_result and self.page_images:
                text_layer = getattr(self.layer_result, 'text_layer', None)
                if text_layer and hasattr(text_layer, 'blocks'):
                    for i, block in enumerate(text_layer.blocks):
                        page_num = getattr(block, 'page_num', 0)
                        if page_num >= len(self.page_images):
                            continue
                        
                        bbox = getattr(block, 'bbox', None)
                        if not bbox or len(bbox) < 4:
                            continue
                        
                        page_img = self.page_images[page_num]
                        try:
                            cropped = page_img.crop(bbox[:4])
                            cropped.save(os.path.join(output_dir, f"text_{i+1:04d}.png"))
                            exported += 1
                        except:
                            pass
                
                # 画像レイヤー
                image_layer = getattr(self.layer_result, 'image_layer', None)
                if image_layer and hasattr(image_layer, 'images'):
                    for i, img_block in enumerate(image_layer.images):
                        if hasattr(img_block, 'image') and img_block.image:
                            img_block.image.save(os.path.join(output_dir, f"image_{i+1:04d}.png"))
                            exported += 1
            
            print(f"✅ PNG exported: {exported} files to {output_dir}")
            return True
            
        except Exception as e:
            print(f"❌ PNG export error: {e}")
            return False
